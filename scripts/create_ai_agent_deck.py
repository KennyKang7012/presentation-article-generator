from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = "outputs/ai_agent_intro_deck.pptx"

WIDE = (13.333, 7.5)
COLORS = {
    "charcoal": "1F2428",
    "ink": "25313A",
    "teal": "0B7C83",
    "mint": "39B7A5",
    "coral": "E85D4F",
    "gold": "F2C14E",
    "cream": "F7F3EA",
    "white": "FFFFFF",
    "muted": "64727D",
    "line": "CBD6D6",
    "pale": "E8F2EF",
    "lilac": "DDD7FB",
    "amber": "F5A623",
    "aqua": "CDEFE9",
    "pink": "F5C7EC",
    "violet": "5A56C8",
    "deepviolet": "373076",
    "green": "009B7A",
}


def rgb(hex_color):
    hex_color = hex_color.strip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def add_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def add_text(slide, text, x, y, w, h, size=18, color="ink", bold=False, font="Aptos", align=None, valign=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    if valign:
        tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(COLORS[color] if color in COLORS else color)
    return box


def add_title(slide, title, subtitle=None, dark=False):
    add_text(slide, title, 0.65, 0.38, 8.8, 0.56, size=27, color="white" if dark else "ink", bold=True, font="Aptos Display")
    if subtitle:
        add_text(slide, subtitle, 0.67, 0.94, 8.2, 0.28, size=8.5, color="mint" if dark else "teal", bold=True)
    add_text(slide, "AI AGENT", 11.45, 0.45, 1.2, 0.18, size=7.5, color="muted" if not dark else "cream", bold=True, align=PP_ALIGN.RIGHT)


def add_footer(slide, dark=False):
    add_text(
        slide,
        "整理自 vocus 文章：AI Agent 入門：解構核心原理與框架",
        0.65,
        7.08,
        7.2,
        0.18,
        size=7.2,
        color="muted" if not dark else "cream",
    )


def rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(COLORS[fill] if fill in COLORS else fill)
    if line:
        shape.line.color.rgb = rgb(COLORS[line] if line in COLORS else line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def circle(slide, x, y, d, fill, line=None):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(COLORS[fill] if fill in COLORS else fill)
    if line:
        shape.line.color.rgb = rgb(COLORS[line] if line in COLORS else line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def line(slide, x1, y1, x2, y2, color="line", width=1.5):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = rgb(COLORS[color] if color in COLORS else color)
    conn.line.width = Pt(width)
    return conn


def label_in_shape(slide, text, x, y, w, h, size=13, color="ink", bold=True):
    box = add_text(slide, text, x, y, w, h, size=size, color=color, bold=bold, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return box


def bullet_list(slide, items, x, y, w, h, color="ink", size=15, gap=0.38):
    for idx, (head, body) in enumerate(items):
        yy = y + idx * gap
        circle(slide, x, yy + 0.06, 0.11, "coral")
        add_text(slide, head, x + 0.22, yy, 1.85, 0.22, size=size, color=color, bold=True)
        add_text(slide, body, x + 2.2, yy, w - 2.2, 0.28, size=size - 1.5, color="cream" if color == "white" else "muted")


def card(slide, title, body, x, y, w, h, accent="teal"):
    rect(slide, x, y, w, h, "white", line="line", radius=True)
    circle(slide, x + 0.25, y + 0.24, 0.28, accent)
    add_text(slide, title, x + 0.65, y + 0.22, w - 0.85, 0.28, size=14.5, color="ink", bold=True)
    add_text(slide, body, x + 0.25, y + 0.76, w - 0.5, h - 0.9, size=11.8, color="muted")


def diagram_box(slide, text, x, y, w, h, fill, size=10.5, color="ink", bold=True):
    rect(slide, x, y, w, h, fill, line=None, radius=True)
    add_text(slide, text, x + 0.08, y + h * 0.38 - 0.06, w - 0.16, 0.24, size=size, color=color, bold=bold, align=PP_ALIGN.CENTER)


def mini_caption(slide, text, x, y, w, fill="pale"):
    rect(slide, x, y, w, 0.22, fill, radius=True)
    add_text(slide, text, x + 0.04, y + 0.06, w - 0.08, 0.08, size=5.8, color="muted", bold=True, align=PP_ALIGN.CENTER)


def draw_architecture_map(slide, x=0.65, y=1.35, scale=1.0):
    def sx(v):
        return x + v * scale

    def sy(v):
        return y + v * scale

    def sw(v):
        return v * scale

    # Central agent and four core modules
    rect(slide, sx(5.05), sy(1.95), sw(2.1), sw(1.65), "deepviolet", radius=True)
    add_text(slide, "AI Agent", sx(5.35), sy(2.62), sw(1.5), sw(0.22), size=18 * scale, color="white", bold=True, align=PP_ALIGN.CENTER)
    diagram_box(slide, "1. Perception", sx(5.08), sy(1.25), sw(2.05), sw(0.46), "amber", size=11 * scale, color="white")
    diagram_box(slide, "2. Planning", sx(2.95), sy(2.52), sw(1.82), sw(0.46), "violet", size=11 * scale, color="white")
    diagram_box(slide, "3. Action", sx(5.08), sy(3.85), sw(2.05), sw(0.46), "coral", size=11 * scale, color="white")
    diagram_box(slide, "4. Memory", sx(7.55), sy(2.52), sw(1.82), sw(0.46), "green", size=11 * scale, color="white")
    line(slide, sx(6.1), sy(1.72), sx(6.1), sy(1.95), "line", 1.5)
    line(slide, sx(4.77), sy(2.75), sx(5.05), sy(2.75), "line", 1.5)
    line(slide, sx(7.15), sy(2.75), sx(7.55), sy(2.75), "line", 1.5)
    line(slide, sx(6.1), sy(3.6), sx(6.1), sy(3.85), "line", 1.5)

    # Perception top row
    for tx, label in [(3.0, "感測器"), (5.25, "處理單元\nGPU / CPU"), (7.65, "多模態\nModel")]:
        diagram_box(slide, label, sx(tx), sy(0.0), sw(1.85), sw(0.58), "gold", size=8.8 * scale)
        line(slide, sx(tx + 0.92), sy(0.58), sx(6.1), sy(1.25), "line", 1)

    # Planning left stack
    planning = [
        ("規劃演算法\nPlanning Algorithm", 0.95, 1.18),
        ("大語言模型\nLLM", 0.95, 2.04),
        ("知識圖譜\nKnowledge Graph", 0.95, 2.9),
        ("推理引擎\nInference Engine", 0.95, 3.76),
    ]
    for label, bx, by in planning:
        diagram_box(slide, label, sx(bx), sy(by), sw(2.0), sw(0.58), "lilac", size=8.3 * scale)
        line(slide, sx(bx + 2.0), sy(by + 0.29), sx(2.95), sy(2.75), "line", 1)
    mini_caption(slide, "複雜任務分解", sx(0.35), sy(1.25), sw(0.45), "lilac")
    mini_caption(slide, "制定行動計劃", sx(0.35), sy(3.86), sw(0.45), "lilac")

    # Memory right stack
    memories = [
        ("原始知識庫\nKM", 9.8, 1.15),
        ("短期記憶", 9.8, 2.42),
        ("長期記憶", 9.8, 3.5),
    ]
    for label, bx, by in memories:
        diagram_box(slide, label, sx(bx), sy(by), sw(2.0), sw(0.62), "aqua", size=8.8 * scale)
        line(slide, sx(9.8), sy(by + 0.31), sx(9.37), sy(2.75), "line", 1)
    line(slide, sx(10.8), sy(3.04), sx(10.8), sy(3.5), "line", 1)
    line(slide, sx(11.05), sy(3.5), sx(11.05), sy(3.04), "line", 1)
    add_text(slide, "構成", sx(10.48), sy(3.12), sw(0.28), sw(0.1), size=5.5 * scale, color="muted", bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "讀取", sx(11.08), sy(3.12), sw(0.28), sw(0.1), size=5.5 * scale, color="muted", bold=True, align=PP_ALIGN.CENTER)

    # Action bottom row
    actions = [
        ("外部工具\nTools", 3.0, 4.86, "API整合、數據分析、時程控管"),
        ("外部服務", 5.2, 4.86, "SQL、Cloud Service、Search"),
        ("反饋與決策機制", 7.45, 4.86, "評估、回饋、下一輪規劃"),
    ]
    for label, bx, by, cap in actions:
        diagram_box(slide, label, sx(bx), sy(by), sw(1.95), sw(0.62), "pink", size=8.5 * scale)
        line(slide, sx(bx + 0.98), sy(by), sx(6.1), sy(4.31), "line", 1)
        mini_caption(slide, cap, sx(bx), sy(by + 0.74), sw(1.95), "pink")


def slide_architecture_overview(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLORS["white"])
    add_title(slide, "AI Agent 參考架構圖", "參考截圖重新繪製：從環境輸入到行動回饋")
    draw_architecture_map(slide, x=0.68, y=1.28, scale=1.0)
    add_footer(slide)
    return slide


def slide_architecture_perception_planning(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLORS["cream"])
    add_title(slide, "感知與規劃：把世界轉成可執行任務", "上半場決定 Agent 看見什麼、理解什麼、打算怎麼做")
    rect(slide, 0.85, 1.55, 5.55, 4.75, "white", line="line", radius=True)
    add_text(slide, "1. Perception", 1.2, 1.92, 2.0, 0.24, size=16, color="amber", bold=True)
    bullet_list(
        slide,
        [
            ("輸入來源", "感測器、文件、圖像、聲音、系統事件。"),
            ("處理單元", "CPU / GPU 負責資料處理與模型推論。"),
            ("多模態模型", "把不同訊號轉成 Agent 可理解的狀態。"),
        ],
        1.2,
        2.55,
        4.65,
        1.6,
        size=12.2,
        gap=0.7,
    )
    rect(slide, 6.95, 1.55, 5.55, 4.75, "charcoal", radius=True)
    add_text(slide, "2. Planning", 7.3, 1.92, 2.0, 0.24, size=16, color="mint", bold=True)
    bullet_list(
        slide,
        [
            ("任務分解", "規劃演算法把目標拆成子任務。"),
            ("推理核心", "LLM、知識圖譜與推理引擎共同產生步驟。"),
            ("計劃輸出", "形成可呼叫工具、可驗證結果的行動方案。"),
        ],
        7.3,
        2.55,
        4.65,
        1.6,
        color="white",
        size=12.2,
        gap=0.7,
    )
    line(slide, 5.85, 3.85, 7.05, 3.85, "coral", 3)
    add_text(slide, "狀態 → 計劃", 5.98, 3.45, 0.95, 0.18, size=9, color="coral", bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)
    return slide


def slide_architecture_action_memory(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLORS["white"])
    add_title(slide, "行動與記憶：讓 Agent 形成閉環", "下半場決定 Agent 做了什麼，以及下一次是否更好")
    rect(slide, 0.9, 1.55, 3.55, 4.85, "pink", radius=True)
    add_text(slide, "3. Action", 1.25, 1.98, 1.6, 0.24, size=16, color="ink", bold=True)
    add_text(slide, "外部工具\nAPI 整合、數據分析、時程控管\n\n外部服務\nSQL、雲端服務、搜尋服務\n\n回饋機制\n評估結果並觸發下一輪決策", 1.25, 2.55, 2.65, 2.35, size=12.8, color="ink")
    rect(slide, 5.0, 1.55, 3.15, 4.85, "charcoal", radius=True)
    add_text(slide, "AI Agent\n執行閉環", 5.58, 3.25, 1.95, 0.5, size=20, color="white", bold=True, align=PP_ALIGN.CENTER)
    rect(slide, 8.7, 1.55, 3.75, 4.85, "aqua", radius=True)
    add_text(slide, "4. Memory", 9.08, 1.98, 1.6, 0.24, size=16, color="green", bold=True)
    add_text(slide, "原始知識庫\n企業資料、文件、規則\n\n短期記憶\n任務上下文與最近互動\n\n長期記憶\n可重用經驗、偏好與策略", 9.08, 2.55, 2.72, 2.35, size=12.8, color="ink")
    line(slide, 4.45, 3.3, 5.0, 3.3, "coral", 2.4)
    line(slide, 8.15, 3.3, 8.7, 3.3, "green", 2.4)
    line(slide, 6.55, 5.95, 6.55, 6.55, "teal", 2)
    add_text(slide, "治理重點：權限、審計、人工覆核、成本監控", 3.55, 6.6, 6.0, 0.18, size=10.5, color="teal", bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)
    return slide


def slide1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLORS["charcoal"])
    for x, y, d, c in [(8.5, 1.0, 0.28, "mint"), (10.8, 1.55, 0.2, "gold"), (9.65, 3.05, 0.36, "coral"), (11.25, 4.25, 0.24, "mint"), (8.15, 4.85, 0.18, "gold")]:
        circle(slide, x, y, d, c)
    for a, b in [((8.64, 1.14), (10.9, 1.65)), ((10.9, 1.65), (9.83, 3.23)), ((9.83, 3.23), (11.37, 4.37)), ((9.83, 3.23), (8.24, 4.94)), ((8.64, 1.14), (9.83, 3.23))]:
        line(slide, a[0], a[1], b[0], b[1], "mint", 1)
    add_text(slide, "AI Agent 入門", 0.75, 1.3, 6.9, 0.7, size=40, color="white", bold=True, font="Aptos Display")
    add_text(slide, "從核心原理、開發框架到負責任導入", 0.8, 2.18, 5.8, 0.35, size=18, color="cream")
    rect(slide, 0.8, 3.05, 4.55, 1.0, "teal", radius=True)
    add_text(slide, "能感知、會規劃、可行動、懂記憶的 AI 系統", 1.12, 3.35, 3.95, 0.32, size=15, color="white", bold=True)
    add_text(slide, "整理自 vocus 文章｜2025/02/24 發佈，2025/08/13 更新", 0.82, 6.68, 5.8, 0.18, size=8.5, color="cream")
    return slide


def slide2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLORS["cream"])
    add_title(slide, "為什麼現在談 AI Agent", "從生成內容，走向自主完成任務")
    rect(slide, 0.8, 1.65, 3.25, 3.6, "charcoal", radius=True)
    add_text(slide, "2025-2026", 1.12, 2.22, 2.6, 0.55, size=33, color="gold", bold=True, font="Aptos Display")
    add_text(slide, "文章引用市場觀點，認為這段期間將是 AI 代理廣泛應用的重要窗口。", 1.14, 3.05, 2.45, 0.82, size=14, color="white")
    card(slide, "效率提升", "協助處理大量、重複、可拆解的工作。", 4.5, 1.65, 2.55, 1.55, "teal")
    card(slide, "工作重組", "讓人更專注於創造性與策略性思考。", 7.28, 1.65, 2.55, 1.55, "coral")
    card(slide, "自主執行", "從被動回答，轉為依目標規劃與行動。", 10.06, 1.65, 2.55, 1.55, "gold")
    line(slide, 4.7, 4.45, 11.75, 4.45, "line", 2)
    for x, t in [(4.7, "工具"), (6.95, "流程"), (9.2, "治理"), (11.45, "價值")]:
        circle(slide, x - 0.12, 4.33, 0.24, "mint")
        add_text(slide, t, x - 0.35, 4.78, 0.7, 0.2, size=10, color="ink", bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)
    return slide


def slide3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLORS["white"])
    add_title(slide, "AI Agent 是什麼", "不是只回覆指令，而是能閉環完成任務")
    rect(slide, 0.85, 1.55, 5.1, 4.6, "pale", radius=True)
    add_text(slide, "一般 AI 工具", 1.22, 1.95, 2.0, 0.28, size=17, color="ink", bold=True)
    add_text(slide, "被動接收問題\n產生內容或答案\n需要人持續拆解步驟", 1.24, 2.55, 3.9, 1.25, size=17, color="muted")
    line(slide, 5.15, 3.88, 6.95, 3.88, "coral", 4)
    add_text(slide, "升級為", 5.52, 3.48, 0.95, 0.22, size=11, color="coral", bold=True, align=PP_ALIGN.CENTER)
    rect(slide, 7.15, 1.55, 5.1, 4.6, "charcoal", radius=True)
    add_text(slide, "AI Agent", 7.55, 1.95, 2.0, 0.28, size=17, color="white", bold=True)
    add_text(slide, "感知環境\n自主規劃\n調用工具行動\n透過記憶改善表現", 7.58, 2.55, 3.9, 1.55, size=17, color="cream")
    circle(slide, 11.1, 4.45, 0.55, "mint")
    add_text(slide, "Loop", 11.02, 4.63, 0.72, 0.16, size=9, color="charcoal", bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)
    return slide


def slide4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLORS["cream"])
    add_title(slide, "四大核心能力", "Agent 行為能力的最小集合")
    items = [
        ("Autonomous", "自主決策，降低人工逐步介入", "teal"),
        ("Perception", "理解文字、圖像、聲音或系統狀態", "mint"),
        ("Planning", "把目標拆成可執行步驟", "gold"),
        ("Self-Improving", "從互動、回饋與記憶中優化", "coral"),
    ]
    positions = [(0.9, 1.55), (6.9, 1.55), (0.9, 4.1), (6.9, 4.1)]
    for idx, ((title, body, c), (x, y)) in enumerate(zip(items, positions), 1):
        rect(slide, x, y, 5.25, 1.85, "white", line="line", radius=True)
        circle(slide, x + 0.35, y + 0.35, 0.72, c)
        label_in_shape(slide, f"{idx}", x + 0.35, y + 0.5, 0.72, 0.22, size=15, color="white")
        add_text(slide, title, x + 1.28, y + 0.34, 3.7, 0.28, size=18, color="ink", bold=True)
        add_text(slide, body, x + 1.28, y + 0.88, 3.45, 0.42, size=14.5, color="muted")
    add_footer(slide)
    return slide


def slide5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLORS["white"])
    add_title(slide, "核心架構迴路", "Perception → Planning → Action → Memory")
    cx, cy = 6.65, 3.75
    nodes = [
        ("Perception", "接收環境資訊", 2.0, 1.75, "teal"),
        ("Planning", "分析情境、拆解任務", 8.3, 1.75, "gold"),
        ("Action", "呼叫工具、完成行動", 8.3, 4.65, "coral"),
        ("Memory", "保存經驗與上下文", 2.0, 4.65, "mint"),
    ]
    for x1, y1, x2, y2 in [(4.75, 2.5, 8.3, 2.5), (9.7, 3.0, 9.7, 4.65), (8.3, 5.4, 4.75, 5.4), (3.4, 4.65, 3.4, 3.0)]:
        line(slide, x1, y1, x2, y2, "line", 2.4)
    circle(slide, cx - 0.7, cy - 0.7, 1.4, "charcoal")
    add_text(slide, "LLM\nOrchestrator", cx - 0.5, cy - 0.2, 1.0, 0.38, size=12, color="white", bold=True, align=PP_ALIGN.CENTER)
    for title, body, x, y, c in nodes:
        rect(slide, x, y, 2.75, 1.0, c, radius=True)
        add_text(slide, title, x + 0.22, y + 0.22, 2.2, 0.22, size=14, color="white", bold=True)
        add_text(slide, body, x + 0.22, y + 0.58, 2.2, 0.18, size=8.8, color="white")
    add_footer(slide)
    return slide


def slide6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLORS["cream"])
    add_title(slide, "案例：餐廳服務生 Agent", "用日常場景理解抽象架構")
    steps = [
        ("01", "觀察顧客", "辨識一對約二、三十歲情侶，掌握情境。", "Perception"),
        ("02", "規劃推薦", "判斷適合餐點，計算如何推薦高單價組合。", "Planning"),
        ("03", "採取行動", "翻開菜單，介紹情人套餐與餐點特色。", "Action"),
        ("04", "運用記憶", "根據銷售經驗，詢問是否加點紅酒。", "Memory"),
    ]
    x = 0.85
    for idx, (num, title, body, tag) in enumerate(steps):
        xx = x + idx * 3.05
        rect(slide, xx, 1.8, 2.55, 3.55, "white", line="line", radius=True)
        add_text(slide, num, xx + 0.28, 2.1, 0.65, 0.28, size=19, color="coral", bold=True)
        add_text(slide, title, xx + 0.28, 2.72, 1.8, 0.25, size=16, color="ink", bold=True)
        add_text(slide, body, xx + 0.28, 3.25, 1.95, 0.8, size=12.3, color="muted")
        rect(slide, xx + 0.28, 4.65, 1.25, 0.36, "pale", radius=True)
        add_text(slide, tag, xx + 0.42, 4.76, 0.95, 0.12, size=7.3, color="teal", bold=True, align=PP_ALIGN.CENTER)
        if idx < 3:
            line(slide, xx + 2.55, 3.55, xx + 3.0, 3.55, "teal", 2)
    add_footer(slide)
    return slide


def slide7(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLORS["white"])
    add_title(slide, "框架選型與開發流程", "框架是基石，流程與治理才決定能否落地")
    rect(slide, 0.82, 1.55, 3.25, 1.45, "charcoal", radius=True)
    add_text(slide, "LangChain", 1.1, 1.93, 2.25, 0.25, size=18, color="white", bold=True)
    add_text(slide, "模組化、多語言支援，適合快速開發與系統整合。", 1.1, 2.35, 2.42, 0.32, size=10.5, color="cream")
    rect(slide, 0.82, 3.35, 3.25, 1.45, "teal", radius=True)
    add_text(slide, "AutoGen", 1.1, 3.73, 2.25, 0.25, size=18, color="white", bold=True)
    add_text(slide, "主打多智能體協作，適合複雜任務與角色分工。", 1.1, 4.15, 2.42, 0.32, size=10.5, color="white")
    flow = ["定義任務", "選擇框架", "設計架構", "整合工具", "訓練評估", "部署監控"]
    for idx, txt in enumerate(flow):
        x = 4.75 + (idx % 3) * 2.45
        y = 1.65 + (idx // 3) * 2.0
        circle(slide, x, y, 0.52, "gold" if idx < 3 else "coral")
        add_text(slide, str(idx + 1), x + 0.18, y + 0.17, 0.18, 0.12, size=9, color="charcoal", bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, txt, x + 0.72, y + 0.08, 1.35, 0.22, size=14, color="ink", bold=True)
        add_text(slide, "把需求轉為可驗證的工程步驟", x + 0.72, y + 0.45, 1.6, 0.28, size=8.2, color="muted")
    add_footer(slide)
    return slide


def slide8(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLORS["cream"])
    add_title(slide, "導入挑戰", "真正困難的部分通常不在模型展示，而在長期營運")
    issues = [
        ("信任與安全", "資料外洩、權限控管、輸出可信度與責任歸屬。", "coral"),
        ("資料與算力", "優質訓練資料稀缺，運算資源與成本壓力升高。", "teal"),
        ("環境成本", "大型 AI 系統會消耗能源與冷卻水資源。", "gold"),
        ("維運成本", "開發與持續監控成本高，未必比人工便宜。", "mint"),
    ]
    for i, (t, b, c) in enumerate(issues):
        x = 0.9 + (i % 2) * 6.05
        y = 1.65 + (i // 2) * 2.25
        rect(slide, x, y, 5.25, 1.55, "white", line="line", radius=True)
        circle(slide, x + 0.32, y + 0.38, 0.55, c)
        add_text(slide, "!", x + 0.53, y + 0.51, 0.12, 0.14, size=12, color="white", bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, t, x + 1.1, y + 0.34, 2.1, 0.23, size=15.8, color="ink", bold=True)
        add_text(slide, b, x + 1.1, y + 0.82, 3.55, 0.3, size=11.5, color="muted")
    rect(slide, 2.65, 6.02, 8.05, 0.48, "charcoal", radius=True)
    add_text(slide, "短期最適合導入：大量、重複、規則可拆解，且風險可控的任務", 3.2, 6.18, 6.95, 0.12, size=11, color="white", bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)
    return slide


def slide9(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLORS["white"])
    add_title(slide, "學習指南", "從看懂 loop，到能設計可控的 Agent")
    steps = [
        ("概念層", "分清 Agent、LLM、工具、記憶與規劃。"),
        ("架構層", "畫出感知、規劃、行動、記憶的資料流。"),
        ("實作層", "從文章摘要、資料查詢等小任務建立最小可行 Agent。"),
        ("評估層", "檢查成功率、錯誤恢復、資料安全與執行成本。"),
        ("治理層", "加入權限、審計、人工覆核與監控告警。"),
    ]
    for i, (t, b) in enumerate(steps):
        y = 1.55 + i * 0.92
        circle(slide, 1.0, y + 0.08, 0.38, "teal" if i < 2 else "coral" if i == 3 else "gold")
        add_text(slide, str(i + 1), 1.14, y + 0.2, 0.1, 0.1, size=7.8, color="white", bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, t, 1.62, y + 0.05, 1.0, 0.18, size=13, color="ink", bold=True)
        add_text(slide, b, 2.62, y + 0.03, 7.8, 0.24, size=12.3, color="muted")
    rect(slide, 9.95, 1.65, 2.35, 4.2, "pale", radius=True)
    add_text(slide, "練習題", 10.25, 2.0, 1.4, 0.25, size=16, color="teal", bold=True)
    add_text(slide, "1. 重畫餐廳服務生流程\n2. 設計客服 Agent 工具清單\n3. 比較 LangChain / AutoGen\n4. 建立導入風險檢核表", 10.25, 2.55, 1.7, 1.75, size=11, color="ink")
    add_footer(slide)
    return slide


def slide10(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, COLORS["charcoal"])
    add_title(slide, "結論", "Agent 的價值在工作流，而不只是模型能力", dark=True)
    add_text(slide, "AI Agent 是「能做事的 AI 系統」", 0.95, 1.75, 6.9, 0.5, size=29, color="white", bold=True, font="Aptos Display")
    bullet_list(
        slide,
        [
            ("不是單純問答", "它需要感知、規劃、工具行動與記憶。"),
            ("不是自動化萬靈丹", "成本、資料、安全與監控會決定是否值得導入。"),
            ("競爭力在設計", "人與 Agent 的協作流程，將成為下一階段差異化能力。"),
        ],
        1.02,
        2.8,
        6.6,
        1.4,
        color="white",
        size=15,
        gap=0.62,
    )
    for i, (t, c) in enumerate([("感知", "teal"), ("規劃", "gold"), ("行動", "coral"), ("記憶", "mint")]):
        angle_x = 8.15 + (i % 2) * 2.1
        angle_y = 2.0 + (i // 2) * 1.85
        circle(slide, angle_x, angle_y, 1.15, c)
        add_text(slide, t, angle_x + 0.35, angle_y + 0.45, 0.45, 0.16, size=12, color="white", bold=True, align=PP_ALIGN.CENTER)
    line(slide, 8.72, 2.58, 10.25, 2.58, "cream", 1.2)
    line(slide, 10.82, 3.15, 10.82, 3.88, "cream", 1.2)
    line(slide, 10.25, 4.42, 8.72, 4.42, "cream", 1.2)
    line(slide, 8.72, 3.88, 8.72, 3.15, "cream", 1.2)
    add_footer(slide, dark=True)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(WIDE[0])
    prs.slide_height = Inches(WIDE[1])
    for fn in [
        slide1,
        slide2,
        slide3,
        slide4,
        slide5,
        slide_architecture_overview,
        slide_architecture_perception_planning,
        slide_architecture_action_memory,
        slide6,
        slide7,
        slide8,
        slide9,
        slide10,
    ]:
        fn(prs)
    prs.core_properties.title = "AI Agent 入門：從核心原理、開發框架到負責任導入"
    prs.core_properties.subject = "vocus 文章重點整理、簡報大綱與學習指南"
    prs.core_properties.author = "Codex"
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
