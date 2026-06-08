from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import json

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


WIDE = (13.333, 7.5)


DEFAULT_THEME = {
    "fonts": {"heading": "Aptos Display", "body": "Aptos"},
    "colors": {
        "charcoal": "1F2428",
        "ink": "25313A",
        "teal": "0B7C83",
        "mint": "39B7A5",
        "coral": "E85D4F",
        "gold": "F2C14E",
        "cream": "F7F3EA",
        "white": "FFFFFF",
        "muted": "64727D",
        "line": "CBD6D6",
        "pale": "E8F2EF",
    },
}


@dataclass
class DeckTheme:
    colors: dict
    heading_font: str
    body_font: str


def load_theme(path: str | Path | None = None) -> DeckTheme:
    data = DEFAULT_THEME
    if path and Path(path).exists():
        loaded = json.loads(Path(path).read_text(encoding="utf-8"))
        data = {
            "fonts": {**DEFAULT_THEME["fonts"], **loaded.get("fonts", {})},
            "colors": {**DEFAULT_THEME["colors"], **loaded.get("colors", {})},
        }
    return DeckTheme(
        colors=data["colors"],
        heading_font=data["fonts"]["heading"],
        body_font=data["fonts"]["body"],
    )


def rgb(hex_color: str) -> RGBColor:
    value = hex_color.strip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def color(theme: DeckTheme, key_or_hex: str) -> RGBColor:
    return rgb(theme.colors.get(key_or_hex, key_or_hex))


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(WIDE[0])
    prs.slide_height = Inches(WIDE[1])
    return prs


def blank_slide(prs: Presentation, bg: str, theme: DeckTheme):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(theme, bg)
    return slide


def add_text(
    slide,
    theme: DeckTheme,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 16,
    color_name: str = "ink",
    bold: bool = False,
    font: str | None = None,
    align=None,
    valign=None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    if valign:
        tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font or theme.body_font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(theme, color_name)
    return box


def add_title(slide, theme: DeckTheme, title: str, subtitle: str = "", dark: bool = False):
    add_text(
        slide,
        theme,
        title,
        0.65,
        0.38,
        9.2,
        0.55,
        size=27,
        color_name="white" if dark else "ink",
        bold=True,
        font=theme.heading_font,
    )
    if subtitle:
        add_text(
            slide,
            theme,
            subtitle,
            0.67,
            0.94,
            8.6,
            0.26,
            size=8.8,
            color_name="mint" if dark else "teal",
            bold=True,
        )
    add_text(
        slide,
        theme,
        "ARTICLE DECK",
        11.3,
        0.46,
        1.4,
        0.18,
        size=7.2,
        color_name="cream" if dark else "muted",
        bold=True,
        align=PP_ALIGN.RIGHT,
    )


def add_footer(slide, theme: DeckTheme, source: str = "", dark: bool = False):
    footer = source if source else "Article-to-presentation"
    add_text(slide, theme, footer, 0.65, 7.08, 8.5, 0.18, size=7.2, color_name="cream" if dark else "muted")


def rect(slide, theme: DeckTheme, x: float, y: float, w: float, h: float, fill: str, line: str | None = None, radius: bool = True):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(theme, fill)
    if line:
        shape.line.color.rgb = color(theme, line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def circle(slide, theme: DeckTheme, x: float, y: float, d: float, fill: str):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(theme, fill)
    shape.line.fill.background()
    return shape


def line(slide, theme: DeckTheme, x1: float, y1: float, x2: float, y2: float, line_color: str = "line", width: float = 1.5):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color(theme, line_color)
    conn.line.width = Pt(width)
    return conn


def add_callout(slide, theme: DeckTheme, title: str, body: str, x: float, y: float, w: float, h: float, accent: str = "teal"):
    rect(slide, theme, x, y, w, h, "white", line="line")
    circle(slide, theme, x + 0.25, y + 0.24, 0.26, accent)
    add_text(slide, theme, title, x + 0.65, y + 0.23, w - 0.9, 0.25, size=13.5, bold=True)
    add_text(slide, theme, body, x + 0.25, y + 0.72, w - 0.5, h - 0.85, size=10.8, color_name="muted")


def add_bullets(slide, theme: DeckTheme, bullets: list[str], x: float, y: float, w: float, *, dark: bool = False, size: float = 13.2, gap: float = 0.48):
    for idx, item in enumerate(bullets[:6]):
        yy = y + idx * gap
        circle(slide, theme, x, yy + 0.08, 0.1, "coral")
        add_text(slide, theme, item, x + 0.23, yy, w - 0.25, 0.28, size=size, color_name="cream" if dark else "ink")


def add_center_label(slide, theme: DeckTheme, text: str, x: float, y: float, w: float, h: float, *, size: float = 14, fill: str = "teal", text_color: str = "white"):
    rect(slide, theme, x, y, w, h, fill)
    add_text(
        slide,
        theme,
        text,
        x + 0.08,
        y + h * 0.35,
        w - 0.16,
        0.22,
        size=size,
        color_name=text_color,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
